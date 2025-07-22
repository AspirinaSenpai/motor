# app.py
from flask import Flask, render_template, request, redirect, url_for, send_file, flash, jsonify
from werkzeug.utils import secure_filename
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
from google.oauth2.service_account import Credentials
from google.auth.transport.requests import Request
import gspread
import google.generativeai as genai
from fpdf import FPDF
import os
import time
from datetime import datetime
import PyPDF2
import docx
import tempfile
import uuid
from io import BytesIO
import logging
import json

# Configuração básica do Flask
app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY', 'dev-key-insecure-change-me')

# Configuração de logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Configurações
UPLOAD_FOLDER = os.path.join(tempfile.gettempdir(), 'motor_reports')
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

ALLOWED_EXTENSIONS = {'pdf', 'docx'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max

class PDFRelatorio(FPDF):
    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=15)
        self.add_page()
        self.set_left_margin(15)
        self.set_right_margin(15)
        self.set_font("Helvetica", size=12)

    def header(self):
        data = datetime.now().strftime("Data: %d/%m/%Y - %H:%M:%S")
        self.set_font("Helvetica", 'B', 14)
        self.cell(0, 10, "Relatório Técnico do Motor", ln=True, align='C')
        self.set_font("Helvetica", '', 10)
        self.cell(0, 10, data, ln=True, align='C')
        self.ln(5)

    def add_relatorio(self, texto):
        texto = self.limpar_caracteres_especiais(texto)
        for linha in texto.split("\n"):
            linha = linha.strip()
            if linha:
                self.multi_cell(0, 8, linha)
                self.ln(1)

    def limpar_caracteres_especiais(self, texto):
        substituicoes = {
            '–': '-', '—': '-', '´': "'", '“': '"', '”': '"',
            '‘': "'", '’': "'", '…': '...', '®': '(R)',
            '©': '(C)', '™': '(TM)'
        }
        for orig, sub in substituicoes.items():
            texto = texto.replace(orig, sub)

        try:
            return texto.encode('latin-1', 'ignore').decode('latin-1')
        except:
            return texto

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        try:
            email_origem = request.form.get('email_origem')
            email_destino = request.form.get('email_destino')
            senha_app = request.form.get('senha_app')
            assunto = request.form.get('assunto', 'Relatório Técnico do Motor - IA')
            modelo_motor = request.form.get('modelo_motor')
            corrente_nominal = request.form.get('corrente_nominal')
            tensao_nominal = request.form.get('tensao_nominal')
            tipo_ligacao = request.form.get('tipo_ligacao')
            observacoes = request.form.get('observacoes', '')

            if 'gerar_relatorio' in request.form:
                if not modelo_motor:
                    flash('Por favor, informe o modelo do motor.', 'error')
                    return redirect(url_for('index'))

                relatorio = gerar_relatorio_ia(
                    modelo_motor, corrente_nominal, tensao_nominal,
                    tipo_ligacao, observacoes, request.files.get('manual')
                )

                pdf_buffer = BytesIO()
                criar_pdf(relatorio, pdf_buffer)
                pdf_buffer.seek(0)

                report_id = str(uuid.uuid4())
                pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], f'report_{report_id}.pdf')
                with open(pdf_path, 'wb') as f:
                    f.write(pdf_buffer.getvalue())

                return jsonify({
                    'status': 'success',
                    'message': 'Relatório gerado com sucesso!',
                    'report_id': report_id
                })

            elif 'enviar_email' in request.form:
                if not all([email_origem, email_destino, senha_app, modelo_motor]):
                    flash('Por favor, preencha todos os campos obrigatórios.', 'error')
                    return redirect(url_for('index'))

                relatorio = gerar_relatorio_ia(
                    modelo_motor, corrente_nominal, tensao_nominal,
                    tipo_ligacao, observacoes, request.files.get('manual')
                )

                temp_pdf = os.path.join(app.config['UPLOAD_FOLDER'], f'temp_{uuid.uuid4()}.pdf')
                criar_pdf(relatorio, temp_pdf)

                enviar_email(
                    email_origem, email_destino, senha_app, assunto,
                    modelo_motor, observacoes, temp_pdf
                )

                if os.path.exists(temp_pdf):
                    os.remove(temp_pdf)

                flash('E-mail enviado com sucesso!', 'success')
                return redirect(url_for('index'))

        except Exception as e:
            logger.error(f"Erro no processamento: {str(e)}")
            flash(f'Erro no processamento: {str(e)}', 'error')
            return redirect(url_for('index'))

    return render_template('index.html')

@app.route('/download/<report_id>')
def download_report(report_id):
    try:
        pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], f'report_{report_id}.pdf')
        if not os.path.exists(pdf_path):
            flash('Relatório não encontrado ou expirado.', 'error')
            return redirect(url_for('index'))

        return send_file(
            pdf_path,
            as_attachment=True,
            download_name=f'relatorio_motor_{report_id[:8]}.pdf',
            mimetype='application/pdf'
        )
    except Exception as e:
        logger.error(f"Erro ao baixar relatório: {str(e)}")
        flash('Erro ao baixar relatório.', 'error')
        return redirect(url_for('index'))

def extrair_texto_manual(arquivo):
    try:
        if not arquivo or not allowed_file(arquivo.filename):
            return ""

        filename = secure_filename(arquivo.filename)
        temp_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        arquivo.save(temp_path)
        texto = ""

        try:
            if filename.lower().endswith('.pdf'):
                with open(temp_path, 'rb') as f:
                    leitor = PyPDF2.PdfReader(f)
                    texto = "\n".join([pagina.extract_text() for pagina in leitor.pages if pagina.extract_text()])
            elif filename.lower().endswith('.docx'):
                doc = docx.Document(temp_path)
                texto = "\n".join([para.text for para in doc.paragraphs])
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

        return texto
    except Exception as e:
        logger.error(f"Erro ao extrair texto do manual: {str(e)}")
        return ""

def gerar_relatorio_ia(modelo_motor, corrente_nominal, tensao_nominal, tipo_ligacao, observacoes, manual_file):
    try:
        genai.configure(api_key=os.getenv('GOOGLE_API_KEY'))
        ultimo_dado = ler_planilha()
        dados_texto = "\n".join([f"- {k}: {v}" for k, v in ultimo_dado.items()])

        info_motor = f"""
INFORMAÇÕES DO MOTOR:
- Modelo: {modelo_motor}
- Corrente Nominal: {corrente_nominal} A
- Tensão Nominal: {tensao_nominal} V
- Tipo de Ligação: {tipo_ligacao}
"""

        manual_texto = extrair_texto_manual(manual_file) if manual_file and allowed_file(manual_file.filename) else ""

        prompt = f"""
Você é um técnico especialista em manutenção de motores elétricos.

{info_motor}

DADOS COLETADOS DO MOTOR:
{dados_texto}

{manual_texto if not manual_texto else f"MANUAL/FICHA TÉCNICA:\n{manual_texto}"}

{'OBSERVAÇÕES EXTRAS DO USUÁRIO:\n' + observacoes + '\n\n' if observacoes else ''}

Faça uma análise técnica detalhada considerando todas essas informações.
"""
        model = genai.GenerativeModel("gemini-2.5-pro")
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        logger.error(f"Erro ao gerar relatório IA: {str(e)}")
        raise

def criar_pdf(relatorio, output):
    try:
        pdf = PDFRelatorio()
        pdf.add_relatorio(relatorio)
        pdf.output(output)
    except Exception as e:
        logger.error(f"Erro ao criar PDF: {str(e)}")
        raise

def enviar_email(email_origem, email_destino, senha_app, assunto, modelo_motor, observacoes, pdf_path):
    try:
        if not all([email_origem, email_destino, senha_app]):
            raise ValueError("E-mail remetente, destinatário e senha são obrigatórios")

        if not os.path.exists(pdf_path):
            raise FileNotFoundError(f"Arquivo PDF não encontrado: {pdf_path}")

        msg = MIMEMultipart()
        msg['From'] = email_origem
        msg['To'] = email_destino
        msg['Subject'] = assunto

        corpo = f"""
RELATÓRIO TÉCNICO - {modelo_motor}
Data: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}

Observações:
{observacoes if observacoes else "Nenhuma observação adicional"}
"""
        msg.attach(MIMEText(corpo, 'plain', 'utf-8'))

        with open(pdf_path, 'rb') as f:
            part = MIMEApplication(f.read(), _subtype='pdf')
            part.add_header('Content-Disposition', 'attachment', filename=f'Relatorio_{modelo_motor}.pdf')
            msg.attach(part)

        with smtplib.SMTP('smtp.gmail.com', 587, timeout=10) as server:
            server.ehlo()
            server.starttls()
            server.ehlo()

            try:
                server.login(email_origem, senha_app)
            except smtplib.SMTPAuthenticationError as auth_err:
                # Mensagem de erro clara e sugestiva
                raise Exception(
                    "Falha na autenticação SMTP. Verifique:\n"
                    "1. Se a verificação em duas etapas está ativada na conta.\n"
                    "2. Se você está usando uma senha de app gerada no Google (não a senha normal).\n"
                    "3. Se a senha de app está correta e ativa.\n"
                    "4. Se houve bloqueio de segurança pelo Google (verifique sua conta).\n"
                    f"Erro original: {auth_err}"
                )

            server.send_message(msg)

    except Exception as e:
        app.logger.error(f"Erro no envio de e-mail: {str(e)}")
        raise

@app.route('/limpar', methods=['POST'])
def limpar_campos():
    try:
        logger.info("Solicitação de limpeza recebida")
        return jsonify({
            'status': 'success',
            'message': 'Campos limpos com sucesso!',
            'timestamp': datetime.now().isoformat()
        })
    except Exception as e:
        logger.error(f"Erro ao limpar campos: {str(e)}")
        return jsonify({
            'status': 'error',
            'message': 'Erro ao limpar campos'
        }), 500

def ler_planilha():
    try:
        SCOPES = ['https://www.googleapis.com/auth/spreadsheets.readonly']
        SPREADSHEET_ID = "1vsWF18ozVUx3B296GtQYXncYHsG6ihhod6ViAKF7bR0"

        json_str = os.getenv("GOOGLE_CREDENTIALS_JSON")
        if not json_str:
            raise EnvironmentError("Variável de ambiente GOOGLE_CREDENTIALS_JSON não encontrada.")

        creds_dict = json.loads(json_str)
        creds = Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
        creds.refresh(Request())

        client = gspread.authorize(creds)
        sheet = client.open_by_key(SPREADSHEET_ID).sheet1
        valores = sheet.get_all_values()

        if not valores or len(valores) < 2:
            return {}

        cabecalho = valores[0]
        ult_linha = valores[-1]
        ultimo_dado = dict(zip(cabecalho, ult_linha))

        for chave in ultimo_dado:
            valor = ultimo_dado[chave]
            if isinstance(valor, str) and "," in valor:
                try:
                    ultimo_dado[chave] = float(valor.replace(",", "."))
                except ValueError:
                    pass

        return ultimo_dado
    except Exception as e:
        logger.error(f"Erro ao acessar planilha: {str(e)}")
        return {}

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(debug=False, host='0.0.0.0', port=port)
